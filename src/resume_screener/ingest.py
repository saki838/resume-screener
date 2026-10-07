"""File discovery and raw text extraction (PDF required; DOCX/TXT bonus).

Every reader returns (text, hyperlinks). Readers raise on failure; the caller
(pipeline) catches per file so one bad resume never stops the batch.
"""
from __future__ import annotations

import hashlib
import logging
import re
import unicodedata
from pathlib import Path

log = logging.getLogger(__name__)
for _noisy in ("pypdf", "pdfminer", "pdfplumber"):
    logging.getLogger(_noisy).setLevel(logging.ERROR)  # noisy 'wrong pointing object' warnings

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class UnsupportedFormatError(Exception):
    pass


def discover_files(input_dir: Path) -> list[Path]:
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")
    return sorted(p for p in input_dir.rglob("*") if p.is_file() and not p.name.startswith("."))


def sha256_of(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    # Soft hyphens / U+FFFE artefacts that some PDF exporters insert at hyphenation points.
    text = text.replace("\u00ad", "").replace("\ufffe", "").replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdfium_text(path: Path) -> str:
    import pypdfium2 as pdfium

    pages: list[str] = []
    pdf = pdfium.PdfDocument(str(path))
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            pages.append(textpage.get_text_range())
            textpage.close()
            page.close()
    finally:
        pdf.close()
    return "\n".join(pages)


def _pdfplumber_text(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(str(path)) as doc:
        return "\n".join((p.extract_text(x_tolerance=1.5) or "") for p in doc.pages)


def text_quality(text: str) -> float:
    """Heuristic readability of an extraction: recognised non-empty sections,
    minus a penalty for collapsed words ("BuiltanAI-poweredplatform")."""
    from .parsing import split_sections  # local import: parsing is a pure-text module

    text = normalize_text(text)
    if len(text) < 50:
        return -1.0
    sections = {kind for kind, body in split_sections(text) if len(body) > 40}
    tokens = re.findall(r"[A-Za-z]+", text)
    squashed = sum(len(t) > 20 for t in tokens) / max(1, len(tokens))
    return len(sections) - 100 * squashed


def _read_pdf(path: Path) -> tuple[str, list[str]]:
    # Two extractors, keep the better one. pdfium preserves word spacing on PDFs
    # where pdfplumber collapses words; pdfplumber preserves reading order on some
    # design-tool exports where pdfium scrambles headings away from their content.
    candidates: list[str] = []
    errors: list[str] = []
    for extractor in (_pdfium_text, _pdfplumber_text):
        try:
            candidates.append(extractor(path))
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{extractor.__name__}: {exc}")
    if not candidates:
        raise ValueError("PDF could not be opened: " + "; ".join(errors))
    best = max(candidates, key=text_quality)  # ties keep pdfium (first)
    return best, _pdf_links(path)


def _pdf_links(path: Path) -> list[str]:
    """URIs from link annotations. Many resumes show only the word "GitHub"."""
    links: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        for page in reader.pages:
            for annot in page.get("/Annots") or []:
                try:
                    obj = annot.get_object()
                    action = obj.get("/A")
                    uri = action.get_object().get("/URI") if action else None
                    if uri:
                        links.append(str(uri).strip())
                except Exception:  # noqa: BLE001 - one bad annotation shouldn't drop the rest
                    continue
    except Exception as exc:  # noqa: BLE001
        log.debug("Could not read link annotations from %s: %s", path.name, exc)
    return links


def _read_docx(path: Path) -> tuple[str, list[str]]:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    links = [
        rel.target_ref for rel in document.part.rels.values()
        if "hyperlink" in rel.reltype and rel.is_external
    ]
    return "\n".join(parts), links


def _read_txt(path: Path) -> tuple[str, list[str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):  # Notepad on Windows often saves cp1252
        try:
            return raw.decode(encoding), []
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1"), []  # maps every byte; never raises


READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".txt": _read_txt}


def extract_text(path: Path) -> tuple[str, list[str]]:
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise UnsupportedFormatError(f"Unsupported file type '{path.suffix}'")
    text, links = reader(path)
    return normalize_text(text), links
